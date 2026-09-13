#!/usr/bin/env python3
"""Exporta una única fuente Markdown a DOCX y PDF de revisión.

Dependencias: scripts/requirements-tfm.txt. No consulta servicios ni modifica el
código del producto. El anexo matemático se incorpora desde su Markdown revisado.
"""
import argparse
from hashlib import sha256
from html import escape
import json
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from markdown_it import MarkdownIt
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT/'memoria/borrador-memoria-tfm.md')
    parser.add_argument('--docx', type=Path, default=ROOT/'memoria/Borrador memoria TFM - Quirón.docx')
    parser.add_argument('--pdf', type=Path, default=ROOT/'memoria/Memoria TFM - Quirón.pdf')
    args = parser.parse_args()
    source = args.source.read_text()
    annex = (ROOT/'docs/ESTUDIO_RED_OBRERA.md').read_text()
    # Las etiquetas conservan su texto; los iconos decorativos no son evidencia.
    clean = lambda s: re.sub('[✅📐⚠🔬️]', '', s)
    combined = source+'\n\n# Anexo A. Estudio de alternativas de la red obrera\n\n'+annex
    tokens = MarkdownIt('commonmark').enable('table').parse(clean(combined))
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    for n in range(1,4):
        document.styles[f'Heading {n}'].font.name = 'Calibri'
        document.styles[f'Heading {n}'].font.size = Pt([17,14,12][n-1])
    section.header.paragraphs[0].text = 'Quirón · Lorenzo Juan Santacreu Pascual · Revisión de cierre'
    section.header.paragraphs[0].style = 'Caption'
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr', 'PAGE')
    footer._p.append(field)
    document.core_properties.title = 'Quirón — Memoria del TFM'
    document.core_properties.author = 'Lorenzo Juan Santacreu Pascual'
    document.core_properties.subject = 'Revisión de alcance, implementación y evaluación; datos pendientes marcados'

    font_dir = Path('/usr/share/fonts/liberation')
    if not (font_dir/'LiberationSans-Regular.ttf').is_file():
        raise SystemExit('Faltan fuentes Liberation Sans/Mono para generar el PDF')
    for label, file in [('Body','LiberationSans-Regular.ttf'),('Body-Bold','LiberationSans-Bold.ttf'),
        ('Body-Italic','LiberationSans-Italic.ttf'),('Body-BoldItalic','LiberationSans-BoldItalic.ttf'),('Code','LiberationMono-Regular.ttf')]:
        pdfmetrics.registerFont(TTFont(label,str(font_dir/file)))
    pdfmetrics.registerFontFamily('Body',normal='Body',bold='Body-Bold',italic='Body-Italic',boldItalic='Body-BoldItalic')
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('Text',fontName='Body',fontSize=11,leading=16.5,spaceAfter=9,alignment=TA_JUSTIFY))
    styles.add(ParagraphStyle('Cell',fontName='Body',fontSize=9,leading=12,spaceAfter=2))
    styles.add(ParagraphStyle('CodeBlock',fontName='Code',fontSize=8,leading=10,spaceAfter=10,backColor=colors.HexColor('#f1f3f5')))
    for n in range(1,4):
        styles[f'Heading{n}'].fontName = 'Body-Bold'
        styles[f'Heading{n}'].fontSize = [17,14,12][n-1]
        styles[f'Heading{n}'].leading = [22,19,16][n-1]
        styles[f'Heading{n}'].spaceBefore = 15
        styles[f'Heading{n}'].spaceAfter = 12
    story = []

    def html_inline(children):
        out=[]
        for child in children or []:
            if child.type in ('text','html_inline'):out.append(escape(child.content))
            elif child.type=='code_inline':out.append('<font name="Code">'+escape(child.content)+'</font>')
            elif child.type in ('softbreak','hardbreak'):out.append('<br/>')
            elif child.type=='strong_open':out.append('<b>')
            elif child.type=='strong_close':out.append('</b>')
            elif child.type=='em_open':out.append('<i>')
            elif child.type=='em_close':out.append('</i>')
            elif child.type=='link_open':
                href=child.attrGet('href') or ''
                out.append('<link href="'+escape(href,quote=True)+'">' if href.startswith(('http://','https://')) else '<font>')
            elif child.type=='link_close':
                # El manuscrito utiliza enlaces web; el código/ficheros usa texto.
                opened=next((x for x in reversed(out) if x.startswith(('<link','<font>'))), '')
                out.append('</link>' if opened.startswith('<link') else '</font>')
        return ''.join(out)

    def doc_inline(paragraph,children,prefix=''):
        if prefix:paragraph.add_run(prefix)
        bold=italic=False
        for child in children or []:
            if child.type=='strong_open':bold=True
            elif child.type=='strong_close':bold=False
            elif child.type=='em_open':italic=True
            elif child.type=='em_close':italic=False
            elif child.type in ('text','code_inline','html_inline','softbreak','hardbreak'):
                run=paragraph.add_run('\n' if child.type in ('softbreak','hardbreak') else child.content)
                run.bold,run.italic=bold,italic
                if child.type=='code_inline':run.font.name='Consolas';run.font.size=Pt(9)

    i=0; list_depth=0; prefix=''; table_number=0; heading=''; chapter_seen=False
    while i<len(tokens):
        token=tokens[i]
        if token.type=='heading_open':
            inline=tokens[i+1];level=min(int(token.tag[1]),3);heading=inline.content
            if level==1 and (heading.startswith(('Capítulo ','Anexo A.')) or heading in ('RESUMEN','ABSTRACT','TABLA RESUMEN')):
                document.add_page_break();story.append(PageBreak());chapter_seen=True
            doc_inline(document.add_heading(level=level),inline.children)
            story.append(Paragraph(html_inline(inline.children),styles[f'Heading{level}']))
            i+=3;continue
        if token.type=='table_open':
            rows=[];current=[];j=i+1
            while tokens[j].type!='table_close':
                if tokens[j].type=='tr_open':current=[]
                elif tokens[j].type=='inline':current.append(tokens[j])
                elif tokens[j].type=='tr_close':rows.append(current)
                j+=1
            table=document.add_table(rows=len(rows),cols=len(rows[0]));table.style='Table Grid'
            for row, cells in zip(table.rows,rows):
                for cell, value in zip(row.cells,cells):
                    doc_inline(cell.paragraphs[0],value.children)
                    for run in cell.paragraphs[0].runs:run.font.size=Pt(9)
            for cell in table.rows[0].cells:
                for run in cell.paragraphs[0].runs:run.bold=True
            width=A4[0]-2*70.87
            cells=[[Paragraph(html_inline(t.children),styles['Cell']) for t in row] for row in rows]
            pdf_table=Table(cells,colWidths=[width/len(rows[0])]*len(rows[0]),repeatRows=1,hAlign='LEFT')
            pdf_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8edf4')),
                ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#b8c0ca')),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
                ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story.append(pdf_table);table_number+=1
            caption=f'Tabla {table_number}. {heading}'
            document.add_paragraph(caption,style='Caption')
            story.extend([Spacer(1,5),Paragraph(escape(caption),styles['Cell']),Spacer(1,12)])
            i=j+1;continue
        if token.type in ('bullet_list_open','ordered_list_open'):list_depth+=1
        elif token.type in ('bullet_list_close','ordered_list_close'):list_depth-=1
        elif token.type=='list_item_open':prefix='• '
        elif token.type=='paragraph_open':
            inline=tokens[i+1]
            if inline.type=='inline':
                paragraph=document.add_paragraph()
                if list_depth:paragraph.paragraph_format.left_indent=Cm(.5*list_depth)
                doc_inline(paragraph,inline.children,prefix)
                story.append(Paragraph(escape(prefix)+html_inline(inline.children),styles['Text']))
                prefix='';i+=3;continue
        elif token.type=='fence':
            paragraph=document.add_paragraph()
            run=paragraph.add_run(token.content.rstrip());run.font.name='Consolas';run.font.size=Pt(9)
            code=escape(token.content.rstrip()).replace('\n','<br/>').replace('  ','&#160;&#160;')
            story.append(Paragraph(code,styles['CodeBlock']))
        elif token.type=='hr':story.append(Spacer(1,12))
        i+=1
    args.docx.parent.mkdir(parents=True,exist_ok=True)
    document.save(args.docx)

    def page(canvas,doc):
        canvas.saveState();canvas.setFont('Body',8)
        canvas.setFillColor(colors.HexColor('#596574'))
        canvas.drawString(70.87,A4[1]-42,'Quirón · Lorenzo Juan Santacreu Pascual · Revisión de cierre')
        canvas.drawCentredString(A4[0]/2,35,str(doc.page))
        canvas.restoreState()
    pdf=SimpleDocTemplate(str(args.pdf),pagesize=A4,leftMargin=70.87,rightMargin=70.87,topMargin=70.87,bottomMargin=70.87,
        title='Quirón — Memoria del TFM',author='Lorenzo Juan Santacreu Pascual')
    pdf.build(story,onFirstPage=page,onLaterPages=page)
    report={'source_sha256':sha256(args.source.read_bytes()).hexdigest(),
        'annex_sha256':sha256((ROOT/'docs/ESTUDIO_RED_OBRERA.md').read_bytes()).hexdigest(),
        'outputs':{p.name:sha256(p.read_bytes()).hexdigest() for p in (args.docx,args.pdf)},
        'pending_fields':source.count('[COMPLETAR'),'table_count':table_number,
        'status':'Revisión para el autor; no versión final aprobada'}
    report_path=ROOT/'docs/evidencias/2026-09-10/exportacion.json'
    report_path.parent.mkdir(parents=True,exist_ok=True)
    report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
